import os
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab imports
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

def parse_time_to_minutes(time_str):
    try:
        parts = time_str.split(":")
        return int(parts[0]) * 60 + int(parts[1])
    except Exception:
        return 0

def check_and_fallback_grid_params(filtered_df, grid_details, unique_time_strs, lunch_by_day, is_common_lunch):
    """
    If visual grid parameters are missing, dynamically reconstruct them
    using default timeslots and course configurations for backward compatibility.
    """
    if grid_details is not None and unique_time_strs is not None and lunch_by_day is not None:
        return grid_details, unique_time_strs, lunch_by_day, is_common_lunch
        
    from app import find_lunch_slots_for_week, build_weekly_grid
    
    try:
        timeslots_df = pd.read_csv("data/timeslots.csv")
    except Exception:
        timeslots_df = pd.DataFrame()
        
    course_info = {}
    try:
        courses_csv = pd.read_csv("data/courses.csv")
        for _, row in courses_csv.iterrows():
            c_id = str(row["course_id"]).strip()
            dur = 1.0
            if "duration_hours" in row.index:
                try:
                    dur = float(row["duration_hours"])
                except Exception:
                    dur = 1.0
            course_info[c_id] = {"duration_hours": dur}
    except Exception:
        pass
        
    lunch_by_day, is_common = find_lunch_slots_for_week(filtered_df, timeslots_df, course_info)
    _, grid_details, unique_time_strs = build_weekly_grid(filtered_df, timeslots_df, lunch_by_day)
    return grid_details, unique_time_strs, lunch_by_day, is_common

from openpyxl.cell.cell import MergedCell

def write_cell_block(ws, start_row, start_col, end_row, end_col, value, font, fill, alignment, border):
    """
    Populates values and styles to a range of cells, then merges them safely.
    Values are only written to the top-left cell of the range if it is not a MergedCell.
    Never assigns .value or styles to MergedCell instances.
    """
    top_left = ws.cell(row=start_row, column=start_col)
    if not isinstance(top_left, MergedCell):
        top_left.value = value
        if font:
            top_left.font = font
        if fill:
            top_left.fill = fill
        if alignment:
            top_left.alignment = alignment
        if border:
            top_left.border = border

    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            if r != start_row or c != start_col:
                cell.value = None
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if alignment:
                cell.alignment = alignment
            if border:
                cell.border = border
                
    if end_row > start_row or end_col > start_col:
        ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)



def generate_timetable_excel(filtered_df, grid_df, header_info=None, course_details_df=None, breaks_dict=None, grid_details=None, unique_time_strs=None, lunch_by_day=None, is_common_lunch=True):
    """
    Generate a styled multi-sheet Excel workbook using 30-minute rows.
    Sheet 1: Weekly Timetable (visual grid with merges).
    Sheet 2: Course Details (legend).
    Sheet 3: Raw Timetable Data.
    """
    if header_info is None:
        header_info = {}
    if course_details_df is None:
        course_details_df = pd.DataFrame()
        
    grid_details, unique_time_strs, lunch_by_day, is_common_lunch = check_and_fallback_grid_params(
        filtered_df, grid_details, unique_time_strs, lunch_by_day, is_common_lunch
    )

    wb = Workbook()
    
    # ----------------------------------------------------
    # Sheet 1: Weekly Timetable
    # ----------------------------------------------------
    ws_grid = wb.active
    ws_grid.title = "Weekly Timetable"
    ws_grid.views.sheetView[0].showGridLines = True
    
    # Styling definitions
    font_title = Font(name="Arial", size=14, bold=True, color="1F4E79")
    font_header = Font(name="Arial", size=9, italic=True, color="595959")
    font_table_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_cell = Font(name="Arial", size=8.5, bold=False)
    font_cell_bold = Font(name="Arial", size=8.5, bold=True)
    font_break = Font(name="Arial", size=10, bold=True, color="3F3F3F")
    
    fill_table_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_timeslot = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_lecture = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_break = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thick_border_side = Side(border_style="medium", color="1F4E79")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_hdr = Border(left=thin_border_side, right=thin_border_side, top=thick_border_side, bottom=thick_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    # 1. Title Block
    ws_grid["A1"] = header_info.get("university_name", "University Timetable")
    ws_grid["A1"].font = font_title
    ws_grid["A2"] = f"{header_info.get('faculty_name', '')} | {header_info.get('semester_name', '')}"
    ws_grid["A2"].font = font_header
    ws_grid["A3"] = f"{header_info.get('title_label', '')} - {header_info.get('week_name', '')} ({header_info.get('validity_period', '')})"
    ws_grid["A3"].font = font_header
    
    ws_grid.row_dimensions[1].height = 20
    ws_grid.row_dimensions[2].height = 15
    ws_grid.row_dimensions[3].height = 15
    ws_grid.row_dimensions[4].height = 8
    
    # 2. Add visual grid header
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    headers = ["Time Slot"] + days
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws_grid.cell(row=5, column=col_idx, value=header)
        cell.font = font_table_hdr
        cell.fill = fill_table_hdr
        cell.alignment = align_center
        cell.border = border_hdr
    ws_grid.row_dimensions[5].height = 22
    
    # Write visual grid rows
    current_row = 6
    for t_str in unique_time_strs:
        ws_grid.row_dimensions[current_row].height = 28 # standard slot height
        
        # Timeslot cell
        t_cell = ws_grid.cell(row=current_row, column=1, value=t_str)
        t_cell.font = font_cell_bold
        t_cell.fill = fill_timeslot
        t_cell.alignment = align_center
        t_cell.border = border_cell
        
        # Check for common lunch
        if is_common_lunch and t_str == lunch_by_day.get("Monday"):
            write_cell_block(
                ws=ws_grid,
                start_row=current_row,
                start_col=2,
                end_row=current_row,
                end_col=len(headers),
                value="LUNCH BREAK",
                font=font_break,
                fill=fill_break,
                alignment=align_center,
                border=border_cell
            )
            current_row += 1
            continue
            
        for c_idx, day in enumerate(days, 2):
            cell = grid_details.get((t_str, day), {"rowspan": 1, "text": "", "is_start": True, "is_lunch": False})
            
            if not cell["is_start"]:
                continue
                
            rowspan = cell["rowspan"]
            text = cell["text"]
            
            if cell["is_lunch"]:
                write_cell_block(
                    ws=ws_grid,
                    start_row=current_row,
                    start_col=c_idx,
                    end_row=current_row + rowspan - 1,
                    end_col=c_idx,
                    value="LUNCH BREAK",
                    font=font_break,
                    fill=fill_break,
                    alignment=align_center,
                    border=border_cell
                )
            elif text:
                clean_text = text.replace("**", "").replace("*", "")
                write_cell_block(
                    ws=ws_grid,
                    start_row=current_row,
                    start_col=c_idx,
                    end_row=current_row + rowspan - 1,
                    end_col=c_idx,
                    value=clean_text,
                    font=font_cell,
                    fill=fill_lecture,
                    alignment=align_center,
                    border=border_cell
                )
            else:
                write_cell_block(
                    ws=ws_grid,
                    start_row=current_row,
                    start_col=c_idx,
                    end_row=current_row,
                    end_col=c_idx,
                    value="-",
                    font=font_cell,
                    fill=None,
                    alignment=align_center,
                    border=border_cell
                )
                
        current_row += 1
        
    # Adjust column widths
    ws_grid.column_dimensions["A"].width = 16
    for col_idx in range(2, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        ws_grid.column_dimensions[col_letter].width = 24
        
    # ----------------------------------------------------
    # Sheet 2: Course Details (Legend)
    # ----------------------------------------------------
    ws_legend = wb.create_sheet(title="Course Legend")
    ws_legend.views.sheetView[0].showGridLines = True
    
    ws_legend["A1"] = f"Course Information Legend - {header_info.get('title_label', '')}"
    ws_legend["A1"].font = font_title
    ws_legend.row_dimensions[1].height = 22
    ws_legend.row_dimensions[2].height = 8
    
    legend_headers = ["Course Code", "Subject Name", "Lecturer", "Credits", "Assessment Type", "Course Type", "Programs"]
    for col_idx, header in enumerate(legend_headers, 1):
        cell = ws_legend.cell(row=3, column=col_idx, value=header)
        cell.font = font_table_hdr
        cell.fill = fill_table_hdr
        cell.alignment = align_center
        cell.border = border_hdr
    ws_legend.row_dimensions[3].height = 22
    
    leg_row = 4
    if not course_details_df.empty:
        for _, row in course_details_df.iterrows():
            ws_legend.row_dimensions[leg_row].height = 18
            vals = [
                row.get("Course Code", ""),
                row.get("Subject Name", ""),
                row.get("Lecturer", ""),
                row.get("Credits", ""),
                row.get("Assessment Type", ""),
                row.get("Course Type", ""),
                row.get("Programs", "")
            ]
            for col_idx, val in enumerate(vals, 1):
                cell = ws_legend.cell(row=leg_row, column=col_idx, value=str(val))
                cell.font = font_cell
                cell.alignment = align_left if col_idx in [2, 3, 7] else align_center
                cell.border = border_cell
            leg_row += 1
            
    legend_widths = [15, 30, 25, 10, 18, 15, 25]
    for i, w in enumerate(legend_widths, 1):
        ws_legend.column_dimensions[get_column_letter(i)].width = w
        
    # ----------------------------------------------------
    # Sheet 3: Raw Generated Dataset
    # ----------------------------------------------------
    ws_raw = wb.create_sheet(title="Raw Data")
    ws_raw.views.sheetView[0].showGridLines = True
    
    raw_headers = list(filtered_df.columns)
    for col_idx, header in enumerate(raw_headers, 1):
        cell = ws_raw.cell(row=1, column=col_idx, value=header)
        cell.font = font_table_hdr
        cell.fill = fill_table_hdr
        cell.alignment = align_center
        cell.border = border_hdr
    ws_raw.row_dimensions[1].height = 22
    
    raw_row = 2
    if not filtered_df.empty:
        for _, row in filtered_df.iterrows():
            ws_raw.row_dimensions[raw_row].height = 18
            for col_idx, col_name in enumerate(raw_headers, 1):
                cell = ws_raw.cell(row=raw_row, column=col_idx, value=str(row[col_name]))
                cell.font = font_cell
                cell.alignment = align_left
                cell.border = border_cell
            raw_row += 1
            
    for col_idx, col_name in enumerate(raw_headers, 1):
        max_len = max(
            filtered_df[col_name].astype(str).map(len).max() if not filtered_df.empty else 0,
            len(col_name)
        )
        ws_raw.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)
        
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream.getvalue()

def generate_timetable_pdf(filtered_df, grid_df, header_info=None, course_details_df=None, breaks_dict=None, grid_details=None, unique_time_strs=None, lunch_by_day=None, is_common_lunch=True):
    """
    Generate a professional Landscape A4 PDF timetable.
    """
    if header_info is None:
        header_info = {}
    if course_details_df is None:
        course_details_df = pd.DataFrame()
        
    grid_details, unique_time_strs, lunch_by_day, is_common_lunch = check_and_fallback_grid_params(
        filtered_df, grid_details, unique_time_strs, lunch_by_day, is_common_lunch
    )

    pdf_stream = io.BytesIO()
    
    # 0.5-inch margins to maximize printable region
    doc = SimpleDocTemplate(
        pdf_stream,
        pagesize=landscape(A4),
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    style_univ = ParagraphStyle(
        'UnivHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=14,
        textColor=colors.HexColor('#1F4E79'),
        alignment=0
    )
    
    style_subtitle = ParagraphStyle(
        'SubtitleHeader',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#595959'),
        alignment=0
    )
    
    style_cell = ParagraphStyle(
        'GridCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=6.5,
        leading=8,
        alignment=1,
        textColor=colors.HexColor('#1E3A8A')
    )
    
    style_break = ParagraphStyle(
        'BreakCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#3F3F3F')
    )
    
    style_th = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=10,
        alignment=1,
        textColor=colors.white
    )
    
    style_timeslot = ParagraphStyle(
        'TimeslotCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=6.5,
        leading=8,
        alignment=1,
        textColor=colors.HexColor('#4B5563')
    )
    
    # 1. Header block
    header_left = [
        Paragraph(header_info.get("university_name", "University Name").upper(), style_univ),
        Paragraph(header_info.get("faculty_name", "Faculty / Department"), style_subtitle),
        Paragraph(f"<b>{header_info.get('title_label', 'Timetable')}</b> — {header_info.get('semester_name', '')}", style_subtitle),
    ]
    
    header_right = [
        Paragraph(f"<b>Week:</b> {header_info.get('week_name', 'N/A')}", style_subtitle),
        Paragraph(f"<b>Validity:</b> {header_info.get('validity_period', 'N/A')}", style_subtitle),
        Paragraph(f"<b>Print Date:</b> {pd.Timestamp.now().strftime('%Y-%m-%d')}", style_subtitle),
    ]
    
    hdr_table = Table([[header_left, header_right]], colWidths=[5.5 * inch, 2.5 * inch])
    hdr_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('ALIGN', (1,0), (1,0), 'RIGHT'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(hdr_table)
    story.append(Spacer(1, 8))
    
    # 2. Build Table
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    headers = ["Time Slot"] + days
    
    table_data = []
    table_data.append([Paragraph(h, style_th) for h in headers])
    
    t_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D9D9D9')),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ('LEFTPADDING', (0,0), (-1,-1), 2),
        ('RIGHTPADDING', (0,0), (-1,-1), 2),
    ]
    
    for r_idx, t_str in enumerate(unique_time_strs, 1):
        row_content = []
        row_content.append(Paragraph(t_str, style_timeslot))
        t_styles.append(('BACKGROUND', (0, r_idx), (0, r_idx), colors.HexColor('#F3F4F6')))
        
        # Check common lunch
        if is_common_lunch and t_str == lunch_by_day.get("Monday"):
            row_content.append(Paragraph("LUNCH BREAK", style_break))
            for _ in range(len(days)):
                row_content.append("")
            t_styles.append(('SPAN', (1, r_idx), (len(headers) - 1, r_idx)))
            t_styles.append(('BACKGROUND', (1, r_idx), (len(headers) - 1, r_idx), colors.HexColor('#E2EFDA')))
            table_data.append(row_content)
            continue
            
        for c_idx, day in enumerate(days, 1):
            cell = grid_details.get((t_str, day), {"rowspan": 1, "text": "", "is_start": True, "is_lunch": False})
            
            if cell["is_lunch"]:
                row_content.append(Paragraph("LUNCH BREAK", style_break))
                t_styles.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#E2EFDA')))
            elif not cell["is_start"]:
                row_content.append("")
            elif cell["text"]:
                clean_text = cell["text"].replace("**", "").replace("*", "").replace("\n", "<br/>")
                row_content.append(Paragraph(clean_text, style_cell))
                t_styles.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#EFF6FF')))
                
                rowspan = cell["rowspan"]
                if rowspan > 1:
                    t_styles.append(('SPAN', (c_idx, r_idx), (c_idx, r_idx + rowspan - 1)))
            else:
                row_content.append(Paragraph("-", style_cell))
                
        table_data.append(row_content)
        
    # Printable area: 770 width
    col_widths = [80] + [138] * len(days)
    
    grid_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    grid_table.setStyle(TableStyle(t_styles))
    story.append(grid_table)
    story.append(Spacer(1, 10))
    
    # 3. Course Legend Table
    if not course_details_df.empty:
        legend_story = []
        style_legend_title = ParagraphStyle(
            'LegendTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#1F4E79')
        )
        legend_story.append(Paragraph("COURSE INFORMATION LEGEND", style_legend_title))
        legend_story.append(Spacer(1, 3))
        
        legend_headers = ["Code", "Subject Name", "Lecturer", "Credits", "Assessment", "Type", "Programs"]
        legend_data = [[Paragraph(h, style_th) for h in legend_headers]]
        
        style_leg_cell = ParagraphStyle(
            'LegendCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=6.5,
            leading=8
        )
        style_leg_cell_center = ParagraphStyle(
            'LegendCellCenter',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=6.5,
            leading=8,
            alignment=1
        )
        
        for _, row in course_details_df.iterrows():
            row_vals = [
                Paragraph(str(row.get("Course Code", "")), style_leg_cell_center),
                Paragraph(str(row.get("Subject Name", "")), style_leg_cell),
                Paragraph(str(row.get("Lecturer", "")), style_leg_cell),
                Paragraph(str(row.get("Credits", "")), style_leg_cell_center),
                Paragraph(str(row.get("Assessment Type", "")), style_leg_cell_center),
                Paragraph(str(row.get("Course Type", "")), style_leg_cell_center),
                Paragraph(str(row.get("Programs", "")), style_leg_cell)
            ]
            legend_data.append(row_vals)
            
        leg_widths = [55, 175, 130, 45, 75, 55, 235]
        legend_table = Table(legend_data, colWidths=leg_widths)
        legend_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D9D9D9')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 2),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ]))
        legend_story.append(legend_table)
        story.append(KeepTogether(legend_story))
        
    doc.build(story)
    pdf_stream.seek(0)
    return pdf_stream.getvalue()
